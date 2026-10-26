# -*- coding: utf-8 -*-
"""
生成《模拟即时通信系统实现》课程设计报告。

用法：
    python make_report.py              # 生成三张图 + DOCX
    python make_report.py --assets-only  # 只生成 report_assets 下的图

依赖：python-docx、matplotlib
"""

import re
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch, Polygon, Rectangle

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE = Path(__file__).resolve().parent
ASSETS = BASE / "report_assets"
MD_PATH = BASE / "课程设计报告_魏荣赫_55250124.md"
DOCX_PATH = BASE / "课程设计报告_魏荣赫_55250124.docx"

TITLE = "模拟即时通信系统实现"
SUBTITLE = "——面向对象程序课程设计报告"
CLASS_NAME = "软件一班（552401班）"
STUDENT_ID = "55250124"
STUDENT_NAME = "魏荣赫"
TERM = "2026-2027学年第1学期"  # 如老师要求旧模板年份，改这一行即可

# 图形配色
EDGE = "#3B4A5A"
TEXT = "#222222"
TITLE_COLOR = "#1F3864"
BOX_FILL = "#F7F9FC"
HEAD_FILL = "#DCE6F1"
NOTE_FILL = "#FFF8E1"
NOTE_EDGE = "#D6C27A"
GREEN_FILL = "#E2F0D9"
GREEN_EDGE = "#70AD47"
BLUE_FILL = "#DDEBF7"
BLUE_EDGE = "#4472C4"
YELLOW_FILL = "#FFF2CC"
YELLOW_EDGE = "#BF9000"
PURPLE_FILL = "#E4DFEC"
PURPLE_EDGE = "#7030A0"


# ============================================================
# 绘图工具
# ============================================================

def setup_font():
    """优先使用本机中文字体，避免图中出现方块。"""
    for name in ("Microsoft YaHei", "SimHei", "SimSun"):
        try:
            font_manager.findfont(name, fallback_to_default=False)
            plt.rcParams["font.sans-serif"] = [name, "SimHei", "DejaVu Sans"]
            plt.rcParams["font.family"] = "sans-serif"
            plt.rcParams["axes.unicode_minus"] = False
            return name
        except Exception:
            continue
    plt.rcParams["font.sans-serif"] = ["DejaVu Sans"]
    plt.rcParams["axes.unicode_minus"] = False
    return "DejaVu Sans"


def new_canvas(w, h):
    fig = plt.figure(figsize=(w, h), dpi=220)
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, w)
    ax.set_ylim(0, h)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    return fig, ax


def save_canvas(fig, path):
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=220, facecolor="white")
    plt.close(fig)


def normalize(dx, dy):
    length = (dx * dx + dy * dy) ** 0.5
    if length == 0:
        return 1.0, 0.0
    return dx / length, dy / length


def class_box(ax, x, y, w, h, name, lines, head_fill=HEAD_FILL,
              name_size=8.4, line_size=7.0):
    """画一个 UML 类框：标题栏 + 分隔线 + 若干行成员。"""
    # 阴影
    shadow = FancyBboxPatch(
        (x + 0.035, y - 0.035), w, h,
        boxstyle="round,pad=0.012,rounding_size=0.05",
        linewidth=0, facecolor="#D9D9D9", alpha=0.55, zorder=1
    )
    ax.add_patch(shadow)

    box = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.012,rounding_size=0.05",
        linewidth=0.9, edgecolor=EDGE, facecolor=BOX_FILL, zorder=2
    )
    ax.add_patch(box)

    # 标题栏（用类框裁掉圆角外的部分）
    head_h = 0.235
    head = Rectangle(
        (x, y + h - head_h), w, head_h,
        linewidth=0, facecolor=head_fill, zorder=2.1
    )
    ax.add_patch(head)
    head.set_clip_path(box)

    ax.plot([x + 0.02, x + w - 0.02], [y + h - head_h, y + h - head_h],
            color=EDGE, linewidth=0.55, zorder=2.5)

    ax.text(x + w / 2, y + h - 0.035, name,
            ha="center", va="top", fontsize=name_size, fontweight="bold",
            color=TITLE_COLOR, zorder=5)

    start_y = y + h - head_h - 0.065
    for i, line in enumerate(lines):
        ax.text(x + 0.055, start_y - i * 0.135, line,
                ha="left", va="top", fontsize=line_size, color=TEXT, zorder=5)


def note_box(ax, x, y, w, h, text, fontsize=6.8):
    box = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.012,rounding_size=0.05",
        linewidth=0.7, edgecolor=NOTE_EDGE, facecolor=NOTE_FILL, zorder=2
    )
    ax.add_patch(box)
    ax.text(x + 0.07, y + h - 0.06, text,
            ha="left", va="top", fontsize=fontsize, color="#4A3B00",
            linespacing=1.35, zorder=5)


def arrow(ax, p1, p2, ls="-", color=EDGE, lw=0.9,
          style="-|>", mutation=8, facecolor=None):
    patch = FancyArrowPatch(
        p1, p2, arrowstyle=style, mutation_scale=mutation,
        linewidth=lw, linestyle=ls, color=color,
        shrinkA=0, shrinkB=0, zorder=4
    )
    if facecolor is not None:
        patch.set_facecolor(facecolor)
        patch.set_edgecolor(color)
    ax.add_patch(patch)


def polyline(ax, points, ls="-", color=EDGE, lw=0.9, zorder=3):
    xs = [p[0] for p in points]
    ys = [p[1] for p in points]
    ax.plot(xs, ys, color=color, linewidth=lw, linestyle=ls,
            solid_capstyle="round", solid_joinstyle="round", zorder=zorder)


def diamond_at(ax, point, toward, size=0.075):
    """在 point 处画一个空心菱形，长轴朝向 toward。"""
    dx, dy = normalize(toward[0] - point[0], toward[1] - point[1])
    nx, ny = -dy, dx
    cx, cy = point[0] + dx * size, point[1] + dy * size
    pts = [
        (cx + nx * size * 0.55, cy + ny * size * 0.55),
        (cx + dx * size, cy + dy * size),
        (cx - nx * size * 0.55, cy - ny * size * 0.55),
        (cx - dx * size, cy - dy * size),
    ]
    ax.add_patch(Polygon(pts, closed=True, facecolor="white",
                         edgecolor=EDGE, linewidth=0.75, zorder=6))


def comp_arrow(ax, points, label=None, label_pos=None, label_size=6.4):
    """组合关系：owner 端菱形，part 端箭头。points 从 owner 到 part。"""
    polyline(ax, points)
    diamond_at(ax, points[0], points[1])
    arrow(ax, points[-2], points[-1], mutation=8)
    if label:
        px, py = label_pos if label_pos else (
            (points[0][0] + points[-1][0]) / 2,
            (points[0][1] + points[-1][1]) / 2,
        )
        ax.text(px, py, label, ha="center", va="center", fontsize=label_size,
                color="#555555", zorder=7,
                bbox=dict(boxstyle="round,pad=0.12", fc="white", ec="none", alpha=0.9))


def assoc_arrow(ax, points, label=None, label_pos=None, label_size=6.4, rotation=0):
    """普通关联：浅灰虚线 + 箭头。"""
    color = "#7F7F7F"
    polyline(ax, points, ls=(0, (3, 2)), color=color)
    arrow(ax, points[-2], points[-1], ls="-", mutation=8, color=color)
    if label:
        px, py = label_pos if label_pos else (
            (points[0][0] + points[-1][0]) / 2,
            (points[0][1] + points[-1][1]) / 2,
        )
        ax.text(px, py, label, ha="center", va="center", fontsize=label_size,
                color="#555555", rotation=rotation, zorder=7,
                bbox=dict(boxstyle="round,pad=0.1", fc="white", ec="none", alpha=0.9))


def inheritance(ax, child, parent, label=None):
    """继承：空心三角在 parent 端。"""
    arrow(ax, child, parent, style="-|>", mutation=11,
          facecolor="white", color=EDGE)
    if label:
        ax.text((child[0] + parent[0]) / 2 + 0.06,
                (child[1] + parent[1]) / 2,
                label, ha="left", va="center", fontsize=6.3,
                color="#555555", zorder=7)


# ============================================================
# 三张图
# ============================================================

def draw_inheritance(path):
    fig, ax = new_canvas(6.5, 4.8)

    # ---- 左：BaseInfoWRH -> UserWRH ----
    class_box(ax, 0.15, 3.75, 1.75, 0.72, "BaseInfoWRH", [
        "- nick_ birth_ applyDate_ city_",
        "+ show()  + tAge()",
    ])
    class_box(ax, 0.15, 2.52, 1.75, 0.88, "UserWRH", [
        "- mainNo_ wxNo_ accs_",
        "+ openService() closeService()",
        "+ login() bindWxToQQ()",
    ])
    inheritance(ax, (1.025, 3.40), (1.025, 3.75), "继承")

    # ---- 中：MicroAccountWRH 家族 ----
    class_box(ax, 2.25, 3.62, 1.95, 0.98, "MicroAccountWRH", [
        "- owner_ type_ logged_",
        "- friends_ : FriendBookWRH",
        "+ svcName()* isSharedID()*",
        "+ isWeChat()* svcNo()*",
    ])
    class_box(ax, 2.40, 2.55, 1.80, 0.66, "QQAccountWRH", [
        "+ svcNo() = 主号",
        "+ isSharedID() = true",
    ])
    class_box(ax, 2.40, 1.60, 1.80, 0.66, "WeiboAccountWRH", [
        "+ svcNo() = 主号",
        "+ isSharedID() = true",
    ])
    class_box(ax, 2.40, 0.62, 1.80, 0.82, "WeChatAccountWRH", [
        "- wxNo_ boundQQ_",
        "+ svcNo() = wxNo_",
        "+ bindQQ() unbindQQ()",
    ])

    # 继承总线：从 MicroAccount 左下角出发，向下再横向分叉
    polyline(ax, [(2.35, 3.62), (2.15, 3.62), (2.15, 1.03)])
    for y in (2.88, 1.93, 1.03):
        arrow(ax, (2.15, y), (2.40, y), mutation=8)
    ax.text(2.02, 2.25, "继承", rotation=90, ha="center", va="center",
            fontsize=6.3, color="#555555")

    # ---- 右：GroupWRH 家族 ----
    class_box(ax, 4.45, 3.62, 1.90, 0.98, "GroupWRH", [
        "- gno_ name_ svc_",
        "- members_ : MemberBook",
        "+ join() kick() setAdmin()",
        "+ canKick()* ... modeTwin()*",
    ])
    class_box(ax, 4.60, 2.55, 1.75, 0.78, "QQGroupWRH", [
        "+ canApplyJoin() = true",
        "+ hasAdminRole() = true",
        "+ createSubGroup()",
    ])
    class_box(ax, 4.60, 1.45, 1.75, 0.78, "WeChatGroupWRH", [
        "+ canApplyJoin() = false",
        "+ hasAdminRole() = false",
        "+ 不支持子群",
    ])
    polyline(ax, [(4.55, 3.62), (4.35, 3.62), (4.35, 1.84)])
    for y in (2.94, 1.84):
        arrow(ax, (4.35, y), (4.60, y), mutation=8)
    ax.text(4.22, 2.35, "继承", rotation=90, ha="center", va="center",
            fontsize=6.3, color="#555555")

    # ---- 说明 ----
    note_box(ax, 0.15, 0.45, 1.75, 1.60,
             "说明\n"
             "▲ 空心三角：继承\n"
             "* 纯虚函数\n"
             "BaseInfoWRH 提供\n"
             "昵称/出生/申请/所在地\n"
             "UserWRH 增加主号、\n"
             "微信号、账号数组")
    note_box(ax, 4.45, 0.35, 1.90, 0.80,
             "QQ 群：可申请 / 管理员 /\n"
             "可建临时讨论组\n"
             "微信群：仅邀请 / 仅群主 /\n"
             "无子群", fontsize=6.6)

    save_canvas(fig, path)


def draw_composition(path):
    fig, ax = new_canvas(6.5, 4.8)

    class_box(ax, 2.20, 4.00, 2.10, 0.62, "PlatformWRH", [
        "- users_ : vector<UserWRH*>",
        "- groups_ : vector<GroupWRH*>",
    ])
    class_box(ax, 0.15, 2.95, 1.80, 0.68, "UserWRH", [
        "- mainNo_ wxNo_",
        "- accs_ : vector<Account*>",
    ])
    class_box(ax, 4.55, 2.95, 1.80, 0.68, "GroupWRH", [
        "- gno_ name_ svc_",
        "- members_ : MemberBook",
    ])
    class_box(ax, 0.15, 1.90, 1.80, 0.68, "MicroAccountWRH", [
        "- owner_ type_ logged_",
        "- friends_ : FriendBookWRH",
    ])
    class_box(ax, 2.25, 1.90, 1.80, 0.68, "FriendBookWRH", [
        "- v_ : vector<FriendInfo>",
        "+ add/remove/commonWith()",
    ])
    class_box(ax, 4.55, 1.90, 1.80, 0.68, "GroupMemberBookWRH", [
        "- v_ : vector<GroupMember>",
        "+ add/remove/setRole()",
    ])
    class_box(ax, 2.25, 0.85, 1.80, 0.68, "FriendInfoWRH", [
        "- person_ : UserWRH*",
        "- remark_ : string",
    ])
    class_box(ax, 4.55, 0.85, 1.80, 0.68, "GroupMemberWRH", [
        "- user_ : UserWRH*",
        "- role_ : MemberRole",
    ])

    # Platform -> User / Group
    comp_arrow(ax, [(2.60, 4.00), (2.60, 3.75), (1.05, 3.75), (1.05, 3.63)],
               "1..*", (1.75, 3.82))
    comp_arrow(ax, [(3.90, 4.00), (3.90, 3.75), (5.45, 3.75), (5.45, 3.63)],
               "1..*", (4.70, 3.82))

    # User -> MicroAccount
    comp_arrow(ax, [(1.05, 2.95), (1.05, 2.58)], "1..3", (1.22, 2.76))
    # MicroAccount -> FriendBook
    comp_arrow(ax, [(1.95, 2.24), (2.25, 2.24)], "1", (2.10, 2.33))
    # FriendBook -> FriendInfo
    comp_arrow(ax, [(3.15, 1.90), (3.15, 1.53)], "0..*", (3.32, 1.71))
    # Group -> MemberBook
    comp_arrow(ax, [(5.45, 2.95), (5.45, 2.58)], "1", (5.62, 2.76))
    # MemberBook -> GroupMember
    comp_arrow(ax, [(5.45, 1.90), (5.45, 1.53)], "0..*", (5.62, 1.71))

    # 非拥有关联：FriendInfo -> User（从两列之间的空隙上去）
    assoc_arrow(ax, [(2.25, 1.08), (2.08, 1.08), (2.08, 3.12), (1.95, 3.12)],
                "person_", (2.11, 1.55), rotation=90)
    # 非拥有关联：GroupMember -> User（从最左侧空隙绕上去）
    assoc_arrow(ax, [(4.55, 1.10), (4.35, 1.10), (4.35, 0.45),
                     (0.05, 0.45), (0.05, 3.12), (0.15, 3.12)],
                "user_", (0.09, 1.55), rotation=90)

    # 顶部说明
    ax.text(0.15, 4.70,
            "MyDateWRH 被 BaseInfoWRH / MicroAccountWRH 复用于日期与 T 龄计算",
            ha="left", va="center", fontsize=6.8, color="#4A3B00")
    # 图例
    ax.plot([5.00, 5.22], [4.70, 4.70], color=EDGE, lw=0.9)
    diamond_at(ax, (5.00, 4.70), (5.22, 4.70), size=0.055)
    ax.text(5.28, 4.70, "组合（拥有）", ha="left", va="center",
            fontsize=6.4, color="#555555")
    ax.plot([5.00, 5.22], [4.55, 4.55], color="#7F7F7F", lw=0.9, ls=(0, (3, 2)))
    ax.text(5.28, 4.55, "关联（不拥有）", ha="left", va="center",
            fontsize=6.4, color="#555555")

    save_canvas(fig, path)


def flow_box(ax, x, y, w, h, text, fill=BLUE_FILL, edge=BLUE_EDGE,
             fontsize=7.2, bold=False):
    box = FancyBboxPatch(
        (x, y), w, h,
        boxstyle="round,pad=0.012,rounding_size=0.05",
        linewidth=0.8, edgecolor=edge, facecolor=fill, zorder=2
    )
    ax.add_patch(box)
    ax.text(x + w / 2, y + h / 2, text, ha="center", va="center",
            fontsize=fontsize, color=TEXT, fontweight="bold" if bold else "normal",
            linespacing=1.25, zorder=5)


def flow_diamond(ax, cx, cy, w, h, text, fontsize=7.0):
    pts = [(cx, cy + h / 2), (cx + w / 2, cy), (cx, cy - h / 2), (cx - w / 2, cy)]
    ax.add_patch(Polygon(pts, closed=True, facecolor=YELLOW_FILL,
                         edgecolor=YELLOW_EDGE, linewidth=0.8, zorder=2))
    ax.text(cx, cy, text, ha="center", va="center", fontsize=fontsize,
            color=TEXT, linespacing=1.2, zorder=5)


def draw_flowchart(path):
    fig, ax = new_canvas(6.5, 4.4)

    flow_box(ax, 2.45, 4.00, 1.60, 0.34, "开始", GREEN_FILL, GREEN_EDGE, 7.4, True)
    flow_box(ax, 2.30, 3.50, 1.90, 0.34, "main 设置控制台 UTF-8")
    flow_box(ax, 2.10, 3.00, 2.30, 0.34, "构造 PlatformWRH(file)")
    flow_diamond(ax, 3.25, 2.45, 1.55, 0.58, "load() 读到用户?")

    # startup branch
    arrow(ax, (3.25, 4.00), (3.25, 3.84))
    arrow(ax, (3.25, 3.50), (3.25, 3.34))
    arrow(ax, (3.25, 3.00), (3.25, 2.74))
    flow_box(ax, 0.45, 2.28, 1.25, 0.34, "buildDemo()", YELLOW_FILL, YELLOW_EDGE)
    flow_box(ax, 4.85, 2.28, 1.35, 0.34, "恢复已有数据", YELLOW_FILL, YELLOW_EDGE)
    arrow(ax, (2.48, 2.45), (1.70, 2.45))
    arrow(ax, (4.03, 2.45), (4.85, 2.45))
    ax.text(2.58, 2.55, "否", fontsize=6.6, color="#666666")
    ax.text(4.12, 2.55, "是", fontsize=6.6, color="#666666")

    # merge into menu loop
    flow_box(ax, 2.20, 1.65, 2.00, 0.40, "run() 主菜单循环", PURPLE_FILL, PURPLE_EDGE,
             7.6, True)
    polyline(ax, [(1.08, 2.28), (1.08, 1.85), (2.20, 1.85)])
    arrow(ax, (2.04, 1.85), (2.20, 1.85), mutation=8)
    polyline(ax, [(5.52, 2.28), (5.52, 2.00), (4.20, 2.00)])
    arrow(ax, (4.36, 2.00), (4.20, 2.00), mutation=8)

    # decision
    flow_diamond(ax, 3.20, 1.20, 1.55, 0.55, "输入 0 / EOF?")
    arrow(ax, (3.20, 1.65), (3.20, 1.48))

    # no -> dispatch -> loop
    flow_box(ax, 4.35, 0.90, 1.85, 0.50,
             "分派：自动演示 / 用户 /\n好友 / 群 / 文件", PURPLE_FILL, PURPLE_EDGE, 6.8)
    arrow(ax, (3.98, 1.20), (4.35, 1.20))
    ax.text(4.06, 1.30, "否", fontsize=6.6, color="#666666")
    polyline(ax, [(6.20, 1.15), (6.30, 1.15), (6.30, 1.75), (4.30, 1.75)])
    arrow(ax, (4.30, 1.75), (4.20, 1.75), mutation=8)
    ax.text(6.02, 1.52, "继续", fontsize=6.3, color="#666666")

    # yes -> destructor -> end
    flow_box(ax, 2.15, 0.45, 2.00, 0.38, "析构：save() + clearAll()",
             BLUE_FILL, BLUE_EDGE, 7.0)
    arrow(ax, (3.20, 0.93), (3.20, 0.83))
    flow_box(ax, 2.55, 0.05, 1.20, 0.28, "结束", GREEN_FILL, GREEN_EDGE, 7.2, True)
    arrow(ax, (3.15, 0.45), (3.15, 0.33))

    # note
    ax.text(0.20, 0.22,
            "文件不存在 / 为空 / 没有用户 → buildDemo()",
            ha="left", va="center", fontsize=6.5, color="#666666")

    save_canvas(fig, path)


# ============================================================
# DOCX 工具
# ============================================================

def set_east_asia(run_or_style, name):
    """给 run 或 style 设置中文 eastAsia 字体。"""
    element = run_or_style._element
    if hasattr(element, "get_or_add_rPr"):
        rPr = element.get_or_add_rPr()
    else:
        rPr = element.rPr
    if rPr is None:
        return
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), name)


def set_run_font(run, name="宋体", size=10.5, bold=False, color=None):
    run.font.name = name
    set_east_asia(run, name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def set_table_borders(table, color="BFBFBF", size="4"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    # tblBorders 必须排在 shd / tblLayout / tblCellMar / tblLook 之前
    tblPr.insert_element_before(
        borders, "w:shd", "w:tblLayout", "w:tblCellMar",
        "w:tblLook", "w:tblCaption", "w:tblDescription"
    )


def set_cell_margins(table, top=35, left=70, bottom=35, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.insert_element_before(
        mar, "w:tblLook", "w:tblCaption", "w:tblDescription"
    )


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def add_bottom_border(paragraph, color="1F3864", size="8", space="4"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    # pBdr 必须排在 shd / tabs / spacing / ind / jc 等之前
    pPr.insert_element_before(
        pBdr, "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku",
        "w:wordWrap", "w:overflowPunct", "w:topLinePunct", "w:autoSpaceDE",
        "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid",
        "w:spacing", "w:ind", "w:contextualSpacing", "w:mirrorIndents",
        "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
        "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle",
        "w:rPr", "w:sectPr", "w:pPrChange"
    )


def add_page_number(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, "宋体", 9, color="595959")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    rFonts = OxmlElement("w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rFonts.set(qn(attr), "宋体")
    rPr.append(rFonts)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "595959")
    rPr.append(color)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "18")
    rPr.append(sz)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.text = "1"
    r.append(t)
    fld.append(r)
    paragraph._p.append(fld)
    run2 = paragraph.add_run(" 页")
    set_run_font(run2, "宋体", 9, color="595959")


# ============================================================
# Markdown 解析 + DOCX 生成
# ============================================================

def split_inline(text, base_size=10.5, base_font="宋体", color=None):
    """把 **粗体** 和 `代码` 拆成 run；返回 [(text, kind), ...]。"""
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    out = []
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            out.append((part[2:-2], "bold"))
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            out.append((part[1:-1], "code"))
        else:
            out.append((part, "normal"))
    return out


def add_runs(paragraph, text, base_size=10.5, base_font="宋体", color=None):
    for piece, kind in split_inline(text):
        if kind == "bold":
            run = paragraph.add_run(piece)
            set_run_font(run, base_font, base_size, bold=True, color=color)
        elif kind == "code":
            run = paragraph.add_run(piece)
            set_run_font(run, "Consolas", base_size - 1, color="C7254E")
        else:
            run = paragraph.add_run(piece)
            set_run_font(run, base_font, base_size, color=color)


def add_body(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Pt(21)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(3)
    add_runs(p, text, 10.5)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.74)
    p.paragraph_format.first_line_indent = Cm(-0.37)
    p.paragraph_format.line_spacing = 1.28
    p.paragraph_format.space_after = Pt(1.5)
    run = p.add_run("• ")
    set_run_font(run, "宋体", 10.5)
    add_runs(p, text, 10.5)
    return p


def add_numbered(doc, text, number):
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Cm(0.85)
    p.paragraph_format.first_line_indent = Cm(-0.55)
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(1.5)
    run = p.add_run(f"{number}. ")
    set_run_font(run, "宋体", 10.5, bold=True)
    add_runs(p, text, 10.5)
    return p


def add_heading(doc, text, level):
    if level == 1:
        p = doc.add_paragraph(style="Heading 1")
        add_runs(p, text, 15.0, "黑体", color="000000")
        for run in p.runs:
            run.font.name = "黑体"
            set_east_asia(run, "黑体")
            run.font.size = Pt(15)
            run.bold = True
        add_bottom_border(p, "1F3864", "10", "3")
    elif level == 2:
        p = doc.add_paragraph(style="Heading 2")
        add_runs(p, text, 12.5, "黑体", color="000000")
        for run in p.runs:
            run.font.name = "黑体"
            set_east_asia(run, "黑体")
            run.font.size = Pt(12.5)
            run.bold = True
    else:
        p = doc.add_paragraph(style="Heading 3")
        add_runs(p, text, 11.0, "黑体", color="000000")
        for run in p.runs:
            run.font.name = "黑体"
            set_east_asia(run, "黑体")
            run.font.size = Pt(11)
            run.bold = True
    p.paragraph_format.keep_with_next = True
    return p


def add_image(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(15.8))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(caption)
    set_run_font(r, "宋体", 9, color="595959")
    return p


def add_code(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "D9D9D9", "4")
    set_cell_margins(table, 80, 120, 80, 120)
    table.columns[0].width = Cm(16.0)
    tblW = table._tbl.tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(int(round(16.0 * 567))))
    tblW.set(qn("w:type"), "dxa")
    cell = table.cell(0, 0)
    shade_cell(cell, "F7F7F7")
    cell.width = Cm(16.0)
    first = True
    for line in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_run_font(run, "Consolas", 8.5, color="333333")
    # 代码块后留一点空
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def parse_markdown(md_text):
    lines = md_text.splitlines()
    title = TITLE
    blocks = []
    i = 0
    n = len(lines)
    para_buf = []

    def flush_para():
        nonlocal para_buf
        if para_buf:
            blocks.append(("p", "".join(para_buf).strip()))
            para_buf = []

    while i < n:
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            flush_para()
            i += 1
            continue

        if stripped.startswith("# "):
            flush_para()
            if title == TITLE:
                title = stripped[2:].strip()
            i += 1
            continue

        if stripped.startswith("## "):
            flush_para()
            blocks.append(("h1", stripped[3:].strip()))
            i += 1
            continue

        if stripped.startswith("### "):
            flush_para()
            blocks.append(("h2", stripped[4:].strip()))
            i += 1
            continue

        if stripped.startswith("#### "):
            flush_para()
            blocks.append(("h3", stripped[5:].strip()))
            i += 1
            continue

        if stripped.startswith("!["):
            m = re.match(r"!\[(.*?)\]\((.*?)\)", stripped)
            if m:
                flush_para()
                blocks.append(("img", m.group(1).strip(), m.group(2).strip()))
                i += 1
                continue

        if stripped.startswith("```"):
            flush_para()
            code_lines = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                code_lines.append(lines[i].rstrip("\n"))
                i += 1
            i += 1
            blocks.append(("code", code_lines))
            continue

        if stripped.startswith("|"):
            flush_para()
            table_lines = []
            while i < n and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            if len(table_lines) >= 2:
                rows = []
                for tl in table_lines:
                    cells = [c.strip() for c in tl.strip().strip("|").split("|")]
                    rows.append(cells)
                # 第二行是 |---|---|
                sep = rows[1]
                if all(re.fullmatch(r":?-{3,}:?", c) for c in sep):
                    blocks.append(("table", rows[0], rows[2:]))
                else:
                    blocks.append(("table", rows[0], rows[1:]))
            continue

        if stripped.startswith("- "):
            flush_para()
            blocks.append(("bullet", stripped[2:].strip()))
            i += 1
            continue

        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            flush_para()
            blocks.append(("num", m.group(2).strip(), int(m.group(1))))
            i += 1
            continue

        para_buf.append(stripped)
        i += 1

    flush_para()
    return title, blocks


def cover(doc):
    # 封面
    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("吉林大学软件学院")
    set_run_font(r, "黑体", 22, bold=True, color="1F3864")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run("《面向对象程序课程设计》")
    set_run_font(r, "黑体", 20, bold=True, color="000000")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    r = p.add_run("实验报告")
    set_run_font(r, "黑体", 28, bold=True, color="1F3864")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("实验项目：模拟即时通信系统实现")
    set_run_font(r, "宋体", 14, bold=False, color="000000")

    for _ in range(2):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)

    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "808080", "6")
    set_cell_margins(table, 80, 100, 80, 100)
    widths = (Cm(4.2), Cm(8.6))
    table.columns[0].width = widths[0]
    table.columns[1].width = widths[1]
    tblW = table._tbl.tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(int(round(12.8 * 567))))
    tblW.set(qn("w:type"), "dxa")
    for row in table.rows:
        row.height = Cm(1.05)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

    data = [
        ("实验地点", "计算机大楼实验室"),
        ("座位号", "____________"),
        ("班级", CLASS_NAME),
        ("学号", STUDENT_ID),
        ("姓名", STUDENT_NAME),
    ]
    for i, (label, value) in enumerate(data):
        c0 = table.cell(i, 0)
        c1 = table.cell(i, 1)
        shade_cell(c0, "F2F2F2")
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(label)
        set_run_font(r0, "宋体", 12, bold=True)
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(value)
        set_run_font(r1, "宋体", 12)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(26)
    r = p.add_run(TERM)
    set_run_font(r, "宋体", 14, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run("2026 年 10 月")
    set_run_font(r, "宋体", 12, color="404040")


def body_title(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(TITLE)
    set_run_font(r, "黑体", 18, bold=True, color="1F3864")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(SUBTITLE)
    set_run_font(r, "楷体", 12, color="404040")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"班级：{CLASS_NAME}　学号：{STUDENT_ID}　姓名：{STUDENT_NAME}")
    set_run_font(r, "宋体", 10.5, color="404040")
    add_bottom_border(p, "BFBFBF", "6", "6")


def setup(document):
    sec = document.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.4)

    # Normal 样式
    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    set_east_asia(normal, "宋体")
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(3)

    # 标题样式
    heading_specs = [
        ("Heading 1", 15, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 8, 3),
    ]
    for name, size, before, after in heading_specs:
        st = document.styles[name]
        st.font.name = "黑体"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor(0, 0, 0)
        set_east_asia(st, "黑体")
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def add_body_section(document):
    body = document.add_section(WD_SECTION.NEW_PAGE)
    body.page_width = Cm(21)
    body.page_height = Cm(29.7)
    body.top_margin = Cm(2.2)
    body.bottom_margin = Cm(2.2)
    body.left_margin = Cm(2.4)
    body.right_margin = Cm(2.4)

    # 页码从正文重新开始（pgNumType 必须排在 cols 等之前）
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    body._sectPr.insert_element_before(
        pgNumType, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote",
        "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter",
        "w:docGrid", "w:printerSettings", "w:sectPrChange"
    )

    # 页眉
    header = body.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hr = hp.add_run("吉林大学软件学院《面向对象程序课程设计》实验报告")
    set_run_font(hr, "宋体", 9, color="808080")
    add_bottom_border(hp, "BFBFBF", "4", "2")

    # 页脚
    footer = body.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)
    return body


def add_table(doc, header, rows):
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "BFBFBF", "4")
    set_cell_margins(table, 45, 70, 45, 70)

    widths_map = {
        2: [5.0, 11.2],
        3: [3.8, 5.2, 7.2],
        4: [3.0, 4.4, 4.6, 4.2],
        5: [2.4, 3.4, 3.6, 3.4, 3.4],
        6: [2.2, 3.0, 3.0, 2.8, 2.6, 2.6],
    }
    widths = widths_map.get(ncols, [16.2 / ncols] * ncols)
    for j, w in enumerate(widths):
        table.columns[j].width = Cm(w)
    tblW = table._tbl.tblPr.find(qn("w:tblW"))
    tblW.set(qn("w:w"), str(int(round(sum(widths) * 567))))
    tblW.set(qn("w:type"), "dxa")

    hdr = table.rows[0]
    set_row_cant_split(hdr)
    set_repeat_header(hdr)
    for j, text in enumerate(header):
        cell = hdr.cells[j]
        cell.width = Cm(widths[j])
        shade_cell(cell, "D9E2F3")
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1.0
        add_runs(p, text, 9.0, "宋体", color="000000")
        for run in p.runs:
            run.bold = True

    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for j in range(ncols):
            cell = row.cells[j]
            cell.width = Cm(widths[j])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.0
            text = row_data[j] if j < len(row_data) else ""
            add_runs(p, text, 9.0, "宋体")

    # 表后留一点空
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    return table


def build_docx(blocks):
    doc = Document()
    setup(doc)
    cover(doc)
    add_body_section(doc)
    body_title(doc)

    numbered = 0
    for block in blocks:
        kind = block[0]
        if kind == "h1":
            text = block[1]
            if text.startswith("二、设计") or text.startswith("三、实现"):
                doc.add_page_break()
            add_heading(doc, text, 1)
            numbered = 0
        elif kind == "h2":
            add_heading(doc, block[1], 2)
            numbered = 0
        elif kind == "h3":
            add_heading(doc, block[1], 3)
            numbered = 0
        elif kind == "p":
            add_body(doc, block[1])
        elif kind == "bullet":
            add_bullet(doc, block[1])
        elif kind == "num":
            add_numbered(doc, block[1], block[2])
        elif kind == "img":
            caption, path = block[1], block[2]
            add_image(doc, BASE / path, caption)
        elif kind == "code":
            add_code(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1], block[2])

    doc.save(DOCX_PATH)
    return doc


def main():
    setup_font()
    ASSETS.mkdir(parents=True, exist_ok=True)
    draw_inheritance(ASSETS / "class_inheritance.png")
    draw_composition(ASSETS / "class_composition.png")
    draw_flowchart(ASSETS / "flowchart.png")
    print(f"[图] {ASSETS / 'class_inheritance.png'}")
    print(f"[图] {ASSETS / 'class_composition.png'}")
    print(f"[图] {ASSETS / 'flowchart.png'}")

    if "--assets-only" in sys.argv:
        return

    md_text = MD_PATH.read_text(encoding="utf-8")
    title, blocks = parse_markdown(md_text)
    # 标题以 Markdown 首行为准；封面项目名固定为同一题目
    global TITLE
    TITLE = title
    build_docx(blocks)
    print(f"[文档] {DOCX_PATH}")

    # 读回检查
    check = Document(str(DOCX_PATH))
    print(f"[检查] 节数={len(check.sections)} 段落={len(check.paragraphs)} "
          f"表格={len(check.tables)} 内嵌图片={len(check.inline_shapes)}")


if __name__ == "__main__":
    main()
